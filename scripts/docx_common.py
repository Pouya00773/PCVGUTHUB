#!/usr/bin/env python3
"""
Gemeinsame DOCX-Formathelfer für die Dokumente dieses Repos.

Bislang lagen set_base_style() und heading() nur in build_docs.py. Mit dem
Lernskript kam ein zweiter Nutzer dazu, deshalb liegen sie jetzt hier — damit
Schriftgrad, Absatzabstände und Überschriftenfarbe über alle erzeugten
Dokumente identisch bleiben.
"""
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AKZENT = RGBColor(0x1F, 0x3B, 0x57)
AKZENT_HEX = "1F3B57"


def set_base_style(doc, size=11):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.paragraph_format.space_after = Pt(2)


def heading(doc, text, size=12, upper=True):
    """Überschrift mit dünner Unterlinie (identisch zu build_docs.py)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.upper() if upper else text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = AKZENT
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), AKZENT_HEX)
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def subheading(doc, text, size=11):
    """Zwischenüberschrift ohne Linie."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = AKZENT
    return p


def _field(paragraph, instr):
    """Fügt ein Word-Feld (z. B. PAGE) in einen Absatz ein."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr_el)
    run._r.append(end)
    return run


def add_page_numbers(doc, prefix=""):
    """Fußzeile 'Seite X von Y' für alle Abschnitte."""
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if prefix:
            r = p.add_run(f"{prefix}  ·  ")
            r.font.size = Pt(8)
            r.font.color.rgb = AKZENT
        r = p.add_run("Seite ")
        r.font.size = Pt(8)
        _field(p, "PAGE").font.size = Pt(8)
        r = p.add_run(" von ")
        r.font.size = Pt(8)
        _field(p, "NUMPAGES").font.size = Pt(8)


def callout(doc, text):
    """Grau hinterlegter Merksatz-Absatz."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Pt(12)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "EEF2F7")
    pPr.append(shd)
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), AKZENT_HEX)
    pbdr.append(left)
    pPr.append(pbdr)
    return p
